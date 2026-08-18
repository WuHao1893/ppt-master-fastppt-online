#!/usr/bin/env python3
"""FastPPT Online editable PPTX export worker.

The default engine authors a lockless ppt-master Quick Generate workspace,
validates its SVG sources, converts them to native DrawingML, and records the
upstream postflight and delivery-check receipts. The old python-pptx exporter is
available only as an explicitly enabled development fallback.
"""

from __future__ import annotations

import argparse
import html
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET


CANVAS_WIDTH = 1280
CANVAS_HEIGHT = 720
FONT_FAMILY = "Microsoft YaHei"
PALETTE = (
    ("#F97316", "#FFF1E8", "#152238"),
    ("#0F766E", "#E6F5F2", "#102A2B"),
    ("#2563EB", "#EAF1FF", "#14233F"),
    ("#C026D3", "#FBEAFA", "#32183A"),
)
SUPPORTED_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp"}
PAGE_ROLES = {
    "cover": "cover",
    "toc": "toc",
    "content": "content",
    "section": "section",
    "ending": "ending",
    "other": "content",
}
PRESENTATION_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
DRAWINGML_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Create an editable FastPPT Online deck.")
    parser.add_argument(
        "--output",
        required=True,
        help="PPTX artifact path inside the project export directory.",
    )
    parser.add_argument(
        "--engine",
        choices=("ppt-master", "legacy"),
        default=os.environ.get("PPTX_EXPORT_ENGINE", "ppt-master"),
        help="Export engine. legacy requires FASTPPT_ALLOW_LEGACY_EXPORT=true.",
    )
    return parser


def _xml_text(value: object) -> str:
    text = str(value or "")
    clean = "".join(
        character
        for character in text
        if (
            ord(character) in (0x9, 0xA, 0xD)
            or 0x20 <= ord(character) <= 0xD7FF
            or 0xE000 <= ord(character) <= 0xFFFD
            or 0x10000 <= ord(character) <= 0x10FFFF
        )
    )
    return html.escape(clean, quote=False)


def _display_units(text: str) -> int:
    return sum(
        2 if unicodedata.east_asian_width(character) in {"W", "F", "A"} else 1
        for character in text
    )


def _wrap_one_line(text: str, max_units: int) -> list[str]:
    source = text.strip()
    if not source:
        return [""]
    output: list[str] = []
    current = ""
    for character in source:
        candidate = current + character
        if current and _display_units(candidate) > max_units:
            split_at = max(current.rfind(" "), current.rfind("\t"))
            if split_at >= max(1, len(current) // 2):
                output.append(current[:split_at].rstrip())
                current = current[split_at + 1 :].lstrip() + character
            else:
                output.append(current.rstrip())
                current = character.lstrip()
        else:
            current = candidate
    if current or not output:
        output.append(current.rstrip())
    return output


def _wrap_text(text: str, max_units: int) -> list[str]:
    output: list[str] = []
    paragraphs = str(text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    for index, paragraph in enumerate(paragraphs):
        if not paragraph.strip():
            if output and output[-1] != "":
                output.append("")
            continue
        output.extend(_wrap_one_line(paragraph, max_units))
        if index < len(paragraphs) - 1 and paragraphs[index + 1].strip():
            output.append("")
    while output and output[-1] == "":
        output.pop()
    return output or [""]


def _fit_text(
    text: str,
    width: int,
    height: int,
    sizes: tuple[int, ...],
    *,
    maximum_lines: int | None = None,
) -> tuple[int, list[str], int]:
    for size in sizes:
        units = max(8, int(width / (size * 0.55)))
        lines = _wrap_text(text, units)
        line_height = max(size + 8, int(size * 1.42))
        if (maximum_lines is None or len(lines) <= maximum_lines) and len(lines) * line_height <= height:
            return size, lines, line_height
    raise ValueError("Page content is too dense for a readable editable slide. Split the page before export.")


def _text_nodes(
    lines: list[str],
    *,
    x: int,
    y: int,
    font_size: int,
    line_height: int,
    color: str,
    weight: str = "normal",
) -> str:
    visible = [(index, line) for index, line in enumerate(lines) if line]
    if not visible:
        return ""
    first_index, _first_line = visible[0]
    attributes = (
        f'x="{x}" y="{y + first_index * line_height}" '
        f'font-family="{FONT_FAMILY}" font-size="{font_size}" '
        f'font-weight="{weight}" fill="{color}"'
    )
    content = ""
    previous_index = first_index
    for position, (index, line) in enumerate(visible):
        dy = 0 if position == 0 else (index - previous_index) * line_height
        content += f'<tspan x="{x}" dy="{dy}">{_xml_text(line)}</tspan>'
        previous_index = index
    return f"<text {attributes}>{content}</text>"


def _page_svg(
    page: dict[str, Any],
    index: int,
    project_name: str,
    visual_filename: str | None,
) -> str:
    accent, soft, ink = PALETTE[index % len(PALETTE)]
    role = PAGE_ROLES.get(str(page.get("pageType", "content")), "content")
    title = str(page.get("title") or "Untitled")
    body = str(page.get("body") or "").strip() or "No body content"
    title_size, title_lines, title_line_height = _fit_text(
        title,
        1120,
        96,
        (44, 40, 36, 32, 28),
        maximum_lines=2,
    )
    body_width = 620 if visual_filename else 760
    body_size, body_lines, body_line_height = _fit_text(
        body,
        body_width,
        350,
        (25, 23, 21, 19, 17, 15),
    )
    title_nodes = _text_nodes(
        title_lines,
        x=64,
        y=122,
        font_size=title_size,
        line_height=title_line_height,
        color=ink,
        weight="700",
    )
    body_nodes = _text_nodes(
        body_lines,
        x=64,
        y=228,
        font_size=body_size,
        line_height=body_line_height,
        color=ink,
    )
    if visual_filename:
        visual = (
            '<g id="visual-anchor" data-pptx-bounds="744 184 478 404">'
            f'<rect x="744" y="184" width="478" height="404" rx="18" fill="#FFFFFF" stroke="{accent}" stroke-width="2"/>'
            f'<image href="../images/{_xml_text(visual_filename)}" x="752" y="192" width="462" height="388" preserveAspectRatio="xMidYMid slice"/>'
            "</g>"
        )
    else:
        visual = (
            '<g id="visual-motif" data-pptx-bounds="878 184 344 404">'
            f'<rect x="878" y="184" width="344" height="404" rx="18" fill="{soft}"/>'
            f'<path d="M916 484 C986 376 1054 430 1178 278" fill="none" stroke="{accent}" stroke-width="8" stroke-linecap="round"/>'
            f'<circle cx="916" cy="484" r="13" fill="{accent}"/>'
            f'<circle cx="1178" cy="278" r="13" fill="{accent}"/>'
            '<rect x="916" y="226" width="112" height="12" rx="6" fill="#64748B"/>'
            '<rect x="916" y="254" width="218" height="8" rx="4" fill="#B9C4D5"/>'
            '<rect x="916" y="274" width="174" height="8" rx="4" fill="#D4DCE7"/>'
            "</g>"
        )
    layout = str(page.get("layout") or "editorial")
    footer_left = f"{index + 1:02d}  /  FASTPPT ONLINE"
    return (
        f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {CANVAS_WIDTH} {CANVAS_HEIGHT}" data-pptx-page-role="{role}">'
        '<rect id="slide-background" data-pptx-role="background" x="0" y="0" width="1280" height="720" fill="#F8FAFC"/>'
        f'<rect id="accent-rail" data-pptx-role="decoration" x="0" y="0" width="18" height="720" fill="{accent}"/>'
        '<g id="header" data-pptx-role="header" data-pptx-bounds="64 44 1152 126">'
        f'<rect x="64" y="46" width="208" height="9" rx="4" fill="{accent}"/>'
        f'<text x="64" y="82" font-family="{FONT_FAMILY}" font-size="14" fill="#64748B">{_xml_text(project_name.upper())}</text>'
        f"{title_nodes}</g>"
        f'<g id="body-copy" data-pptx-bounds="64 194 {body_width} 374">{body_nodes}</g>'
        f"{visual}"
        '<g id="footer" data-pptx-role="footer" data-pptx-bounds="64 664 1152 28">'
        f'<text x="64" y="684" font-family="{FONT_FAMILY}" font-size="12" fill="#64748B">{_xml_text(footer_left)}</text>'
        f'<text x="1216" y="684" font-family="{FONT_FAMILY}" font-size="12" text-anchor="end" fill="#64748B">{_xml_text(layout)}</text>'
        "</g></svg>"
    )


def _copy_visual_asset(
    raw_path: object,
    images_dir: Path,
    page_id: str,
    index: int,
) -> str | None:
    if not raw_path:
        return None
    source = Path(str(raw_path)).expanduser().resolve()
    if not source.is_file():
        raise ValueError(f"Visual asset for {page_id} does not exist.")
    suffix = source.suffix.lower()
    if suffix not in SUPPORTED_IMAGE_SUFFIXES:
        raise ValueError(f"Visual asset for {page_id} has an unsupported extension: {suffix}.")
    filename = f"{index + 1:03d}-{page_id}{suffix}"
    shutil.copy2(source, images_dir / filename)
    return filename


def _prepare_quick_project(project_dir: Path, payload: dict[str, Any]) -> list[dict[str, Any]]:
    svg_dir = project_dir / "svg_output"
    images_dir = project_dir / "images"
    svg_dir.mkdir(parents=True)
    images_dir.mkdir(parents=True)
    pages = list(payload.get("pages") or [])
    if not pages:
        pages = [
            {
                "pageId": "page_empty",
                "pageType": "cover",
                "title": str(payload.get("name") or "FastPPT Online"),
                "body": "Empty project",
                "layout": "editorial",
                "editableLevel": "native_structure",
                "nonEditableRegions": [],
            }
        ]
    reports: list[dict[str, Any]] = []
    width = max(2, len(str(len(pages))))
    for index, page in enumerate(pages):
        page_id = str(page.get("pageId") or f"page_{index + 1}")
        visual_filename = _copy_visual_asset(
            page.get("visualAssetPath"),
            images_dir,
            page_id,
            index,
        )
        svg = _page_svg(page, index, str(payload.get("name") or "FastPPT Online"), visual_filename)
        svg_path = svg_dir / f"{index + 1:0{width}d}_{page_id}.svg"
        svg_path.write_text(svg, encoding="utf-8", newline="\n")
        non_editable = list(page.get("nonEditableRegions") or [])
        if visual_filename and "visual_anchor" not in non_editable:
            non_editable.append("visual_anchor")
        reports.append(
            {
                "page_id": page_id,
                "editable_level": page.get("editableLevel", "native_structure"),
                "non_editable_regions": non_editable,
                "local_image_region": "visual-anchor" if visual_filename else None,
                "source_svg": svg_path.name,
            }
        )
    return reports


def _run_tool(command: list[str], label: str) -> subprocess.CompletedProcess[str]:
    timeout = max(30, int(os.environ.get("PPT_MASTER_TIMEOUT_SECONDS", "180")))
    environment = {**os.environ, "PYTHONUTF8": "1"}
    creation_flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
    try:
        result = subprocess.run(
            command,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout,
            env=environment,
            creationflags=creation_flags,
            check=False,
        )
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError(f"{label} timed out after {timeout} seconds.") from exc
    if result.returncode != 0:
        detail = (result.stderr.strip() or result.stdout.strip())[-6000:]
        raise RuntimeError(f"{label} failed: {detail}")
    return result


def _ppt_master_scripts() -> dict[str, Path]:
    skill_dir = Path(__file__).resolve().parents[3] / "skills" / "ppt-master"
    required = {
        "checker": skill_dir / "scripts" / "svg_quality_checker.py",
        "exporter": skill_dir / "scripts" / "svg_to_pptx.py",
        "delivery": skill_dir / "scripts" / "pptx_delivery_check.py",
        "license": skill_dir / "LICENSE",
        "skill": skill_dir / "SKILL.md",
        "sponsors": skill_dir / "SPONSORS.md",
        "sponsors_cn": skill_dir / "SPONSORS_CN.md",
    }
    missing = [str(path) for path in required.values() if not path.is_file()]
    if missing:
        raise RuntimeError(
            "The official ppt-master distribution is incomplete: " + ", ".join(missing)
        )
    return required


def _load_json(path: Path, label: str) -> dict[str, Any]:
    if not path.is_file():
        raise RuntimeError(f"{label} did not create {path}.")
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, UnicodeError) as exc:
        raise RuntimeError(f"{label} returned invalid UTF-8 JSON: {exc}") from exc
    if not isinstance(value, dict):
        raise RuntimeError(f"{label} JSON must be an object.")
    return value


def _natural_slide_key(name: str) -> tuple[int, str]:
    stem = Path(name).stem
    digits = "".join(character for character in stem if character.isdigit())
    return (int(digits or 0), name)


def _picture_is_full_slide(picture: ET.Element, width: int, height: int) -> bool:
    transform = picture.find(f".//{{{DRAWINGML_NS}}}xfrm")
    if transform is None:
        return False
    offset = transform.find(f"{{{DRAWINGML_NS}}}off")
    extent = transform.find(f"{{{DRAWINGML_NS}}}ext")
    if offset is None or extent is None:
        return False
    try:
        x = int(offset.attrib.get("x", "-1"))
        y = int(offset.attrib.get("y", "-1"))
        cx = int(extent.attrib.get("cx", "0"))
        cy = int(extent.attrib.get("cy", "0"))
    except ValueError:
        return False
    return x <= width * 0.01 and y <= height * 0.01 and cx >= width * 0.99 and cy >= height * 0.99


def _inspect_pptx_structure(output: Path, page_reports: list[dict[str, Any]]) -> dict[str, Any]:
    with zipfile.ZipFile(output) as archive:
        presentation = ET.fromstring(archive.read("ppt/presentation.xml"))
        slide_size = presentation.find(f"{{{PRESENTATION_NS}}}sldSz")
        if slide_size is None:
            raise RuntimeError("PPTX package has no slide-size declaration.")
        slide_width = int(slide_size.attrib["cx"])
        slide_height = int(slide_size.attrib["cy"])
        slide_names = sorted(
            (
                name
                for name in archive.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            ),
            key=_natural_slide_key,
        )
        inspected: list[dict[str, Any]] = []
        for index, slide_name in enumerate(slide_names):
            root = ET.fromstring(archive.read(slide_name))
            text_shapes = sum(
                1
                for shape in root.findall(f".//{{{PRESENTATION_NS}}}sp")
                if any((node.text or "").strip() for node in shape.findall(f".//{{{DRAWINGML_NS}}}t"))
            )
            shape_count = len(root.findall(f".//{{{PRESENTATION_NS}}}sp"))
            connector_count = len(root.findall(f".//{{{PRESENTATION_NS}}}cxnSp"))
            graphic_count = len(root.findall(f".//{{{PRESENTATION_NS}}}graphicFrame"))
            pictures = root.findall(f".//{{{PRESENTATION_NS}}}pic")
            full_slide_picture = any(
                _picture_is_full_slide(picture, slide_width, slide_height)
                for picture in pictures
            )
            report = dict(page_reports[index] if index < len(page_reports) else {})
            report.update(
                {
                    "text_shapes": text_shapes,
                    "native_shapes": shape_count + connector_count + graphic_count,
                    "picture_shapes": len(pictures),
                    "full_slide_raster": full_slide_picture and text_shapes == 0,
                }
            )
            inspected.append(report)
    static_passed = bool(inspected) and all(
        not report["full_slide_raster"]
        and report["text_shapes"] >= 1
        and report["native_shapes"] >= 1
        for report in inspected
    )
    return {
        "page_count": len(inspected),
        "full_slide_raster_count": sum(bool(report["full_slide_raster"]) for report in inspected),
        "static_structure_passed": static_passed,
        "pages": inspected,
    }


def _export_with_ppt_master(
    payload: dict[str, Any],
    output: Path,
) -> dict[str, Any]:
    scripts = _ppt_master_scripts()
    with tempfile.TemporaryDirectory(
        prefix="fastppt-online-export-",
        dir=output.parent,
    ) as temporary_directory:
        project_dir = Path(temporary_directory) / "project"
        project_dir.mkdir()
        page_reports = _prepare_quick_project(project_dir, payload)
        _run_tool(
            [
                sys.executable,
                str(scripts["checker"]),
                str(project_dir),
                "--quick-generate",
                "--format",
                "ppt169",
                "--stage",
                "final",
                "--json",
            ],
            "ppt-master SVG quality checker",
        )
        quality_path = project_dir / "validation" / "svg_quality_report.json"
        quality_report = _load_json(quality_path, "ppt-master SVG quality checker")
        _run_tool(
            [
                sys.executable,
                str(scripts["exporter"]),
                str(project_dir),
                "--quick-generate",
                "--format",
                "ppt169",
                "--no-animations",
                "--no-notes",
                "-o",
                str(output),
            ],
            "ppt-master native DrawingML export",
        )
        postflight_path = project_dir / "validation" / f"{output.stem}.report.json"
        postflight_report = _load_json(postflight_path, "ppt-master postflight")
        delivery_result = _run_tool(
            [sys.executable, str(scripts["delivery"]), str(output)],
            "ppt-master delivery check",
        )
        try:
            delivery_report = json.loads(delivery_result.stdout)
        except json.JSONDecodeError as exc:
            raise RuntimeError(f"ppt-master delivery check returned invalid JSON: {exc}") from exc
        if not isinstance(delivery_report, dict):
            raise RuntimeError("ppt-master delivery check JSON must be an object.")
        structure = _inspect_pptx_structure(output, page_reports)
        if not structure["static_structure_passed"]:
            raise RuntimeError("Generated PPTX failed the editable static-structure gate.")
        return {
            "schema": "fastppt-online.export-qa.v2",
            "project_id": payload.get("projectId"),
            "pptx": output.name,
            "export_engine": "ppt_master_svg_to_drawingml",
            **structure,
            "svg_quality": quality_report,
            "pptx_postflight": postflight_report,
            "delivery_check": delivery_report,
            "powerpoint_render_status": "not_run_by_export_worker",
        }


def _legacy_add_text(
    slide: Any,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    size: int,
    color: str,
    bold: bool = False,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = PP_ALIGN.LEFT
    paragraph.font.name = FONT_FAMILY
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor.from_string(color)


def _export_legacy(payload: dict[str, Any], output: Path) -> dict[str, Any]:
    if os.environ.get("FASTPPT_ALLOW_LEGACY_EXPORT", "").lower() != "true":
        raise RuntimeError(
            "Legacy python-pptx export is a development fallback. Set "
            "FASTPPT_ALLOW_LEGACY_EXPORT=true explicitly to enable it."
        )
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches
    except ImportError as exc:
        raise RuntimeError("Legacy export requires python-pptx.") from exc
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    pages = list(payload.get("pages") or []) or [
        {"pageId": "page_empty", "title": "FastPPT Online", "body": "Empty project"}
    ]
    page_reports: list[dict[str, Any]] = []
    for index, page in enumerate(pages):
        accent, _soft, ink = PALETTE[index % len(PALETTE)]
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("F8FAFC")
        bar = slide.shapes.add_shape(1, 0, 0, Inches(0.19), presentation.slide_height)
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor.from_string(accent.lstrip("#"))
        bar.line.fill.background()
        _legacy_add_text(slide, str(page.get("title") or "Untitled"), 0.65, 1.1, 11.8, 0.8, size=25, color=ink.lstrip("#"), bold=True)
        _legacy_add_text(slide, str(page.get("body") or "No body content"), 0.65, 2.1, 11.8, 4.4, size=15, color=ink.lstrip("#"))
        page_reports.append(
            {
                "page_id": page.get("pageId"),
                "editable_level": page.get("editableLevel", "native_partial"),
                "non_editable_regions": page.get("nonEditableRegions", []),
                "local_image_region": None,
            }
        )
    presentation.save(output)
    structure = _inspect_pptx_structure(output, page_reports)
    return {
        "schema": "fastppt-online.export-qa.v2",
        "project_id": payload.get("projectId"),
        "pptx": output.name,
        "export_engine": "legacy_python_pptx_development_fallback",
        **structure,
        "svg_quality": {"status": "not-run-development-fallback"},
        "pptx_postflight": {"status": "not-run-development-fallback"},
        "delivery_check": {"status": "not-run-development-fallback"},
        "powerpoint_render_status": "not_run_by_export_worker",
    }


def main(argv: list[str] | None = None) -> int:
    sys.stdin.reconfigure(encoding="utf-8", errors="strict")
    sys.stdout.reconfigure(encoding="utf-8", errors="strict")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    args = build_parser().parse_args(argv)
    try:
        payload = json.load(sys.stdin)
    except (json.JSONDecodeError, UnicodeError) as exc:
        print(f"Invalid UTF-8 JSON payload: {exc}", file=sys.stderr)
        return 2
    output = Path(args.output).expanduser().resolve()
    if output.suffix.lower() != ".pptx":
        print("Output path must use the .pptx extension.", file=sys.stderr)
        return 2
    output.parent.mkdir(parents=True, exist_ok=True)
    try:
        qa = (
            _export_legacy(payload, output)
            if args.engine == "legacy"
            else _export_with_ppt_master(payload, output)
        )
        qa_path = output.with_suffix(".qa.json")
        qa_path.write_text(
            json.dumps(qa, ensure_ascii=False, indent=2),
            encoding="utf-8",
            newline="\n",
        )
    except (OSError, RuntimeError, ValueError, zipfile.BadZipFile) as exc:
        print(str(exc), file=sys.stderr)
        return 1
    print(json.dumps({"output": str(output), "qa": str(qa_path)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
