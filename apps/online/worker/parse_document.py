from __future__ import annotations

import json
import os
import re
import sys
import zipfile
from datetime import date, datetime
from pathlib import Path


NUMBER_RE = re.compile(r"(?<![\w.])(?:\d{4}[-/.年]\d{1,2}(?:[-/.月]\d{1,2}日?)?|[-+]?\d[\d,]*(?:\.\d+)?%?)")


def env_limit(name: str, default: int) -> int:
    raw = os.environ.get(name, str(default))
    try:
        value = int(raw)
    except ValueError as error:
        raise ValueError(f"{name} must be an integer") from error
    if value <= 0:
        raise ValueError(f"{name} must be greater than zero")
    return value


def validate_archive(path: Path) -> None:
    entry_limit = env_limit("MAX_DOCUMENT_ARCHIVE_ENTRIES", 5000)
    expanded_limit = env_limit("MAX_DOCUMENT_UNCOMPRESSED_BYTES", 104_857_600)
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > entry_limit:
            raise ValueError(
                f"archive contains {len(entries)} entries; limit is {entry_limit}"
            )
        expanded_size = 0
        for entry in entries:
            if entry.flag_bits & 0x1:
                raise ValueError("encrypted archive entries are not supported")
            expanded_size += entry.file_size
            if expanded_size > expanded_limit:
                raise ValueError(
                    f"archive expands to more than {expanded_limit} bytes"
                )


def property_text(value: object) -> str:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value or "").strip()


def normalize_key(line: str) -> str:
    key = NUMBER_RE.sub("{value}", line.lower())
    key = re.sub(r"\s+", " ", key).strip(" -:：;；,，。")
    return key[:240]


def facts_from_lines(lines: list[str]) -> list[dict[str, str]]:
    facts: list[dict[str, str]] = []
    for index, line in enumerate(lines, start=1):
        compact = " ".join(line.split())
        if not compact:
            continue
        key = normalize_key(compact)
        if not key:
            continue
        for match in NUMBER_RE.finditer(compact):
            value = match.group(0)
            facts.append({
                "key": key,
                "value": value,
                "normalizedValue": value.replace(",", "").lower(),
                "kind": "metric" if "%" in value else "number",
                "location": f"line:{index}",
                "sourceLocator": f"line:{index}",
                "confidence": 0.95,
                "locked": False,
                "context": compact[:500],
            })
    return facts[:2000]


def parse_markdown(path: Path) -> tuple[list[str], list[str], list[str], dict]:
    text = path.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    headings = [line.lstrip("# ").strip() for line in lines if line.lstrip().startswith("#")]
    return lines, headings, [], {}


def parse_docx(path: Path) -> tuple[list[str], list[str], list[str], dict]:
    from docx import Document

    document = Document(path)
    lines: list[str] = []
    headings: list[str] = []
    hyperlink_count = 0
    for paragraph in document.paragraphs:
        hyperlink_count += len(paragraph._p.xpath(".//w:hyperlink"))
        text = "".join(
            str(node.text or "")
            for node in paragraph._p.iter()
            if str(node.tag).endswith("}t")
        ).strip()
        if not text:
            continue
        lines.append(text)
        if paragraph.style and (paragraph.style.name.lower().startswith("heading") or paragraph.style.name.lower() in {"title", "subtitle"}):
            headings.append(text)
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                lines.append(" | ".join(values))
    core = document.core_properties
    properties = {
        key: property_text(getattr(core, key, ""))
        for key in (
            "title",
            "subject",
            "author",
            "keywords",
            "comments",
            "category",
            "identifier",
            "language",
            "last_modified_by",
            "revision",
            "version",
            "created",
            "modified",
        )
        if property_text(getattr(core, key, ""))
    }
    image_count = sum(
        1
        for relationship in document.part.rels.values()
        if relationship.reltype.endswith("/image")
    )
    return lines, headings, [], {
        "paragraphCount": len(document.paragraphs),
        "tableCount": len(document.tables),
        "hyperlinkCount": hyperlink_count,
        "imageCount": image_count,
        "properties": properties,
    }


def parse_pdf(path: Path) -> tuple[list[str], list[str], list[str], dict]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    lines: list[str] = []
    warnings: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        page_lines = [line.strip() for line in text.splitlines() if line.strip()]
        if len("".join(page_lines)) < 20:
            warnings.append(f"第 {index} 页可能是扫描页，文字提取可靠性较低。")
        lines.extend(page_lines)
    headings = [line for line in lines if len(line) <= 60 and not line.endswith(("。", ".", ";", "；"))][:80]
    return lines, headings, warnings, {"pageCount": len(reader.pages)}


def parse_pptx(path: Path) -> tuple[list[str], list[str], list[str], dict]:
    from pptx import Presentation

    presentation = Presentation(path)
    lines: list[str] = []
    headings: list[str] = []
    warnings: list[str] = []
    slides: list[dict] = []
    slide_width = int(presentation.slide_width)
    slide_height = int(presentation.slide_height)
    for index, slide in enumerate(presentation.slides, start=1):
        slide_lines: list[str] = []
        text_shape_count = 0
        image_count = 0
        table_count = 0
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text_shape_count += 1
                text = shape.text.strip()
                if text:
                    slide_lines.extend(part.strip() for part in text.splitlines() if part.strip())
            if getattr(shape, "shape_type", None) == 13:
                image_count += 1
            if getattr(shape, "has_table", False):
                table_count += 1
        title = slide_lines[0] if slide_lines else f"第 {index} 页"
        body = "\n".join(slide_lines[1:])
        if slide_lines:
            headings.append(slide_lines[0])
            lines.extend(slide_lines)
        else:
            warnings.append(f"第 {index} 页没有可提取的原生文本，可能包含不可编辑视觉区域。")
        non_editable = [f"slide:{index}:image-region"] if image_count else []
        slides.append({
            "index": index,
            "title": title,
            "body": body,
            "layoutName": getattr(slide.slide_layout, "name", "unknown") or "unknown",
            "shapeCount": len(slide.shapes),
            "textShapeCount": text_shape_count,
            "imageCount": image_count,
            "tableCount": table_count,
            "editableLevel": "native_partial" if non_editable else "native_structure",
            "nonEditableRegions": non_editable,
        })
    return lines, headings, warnings, {"pageCount": len(slides), "slideWidth": slide_width, "slideHeight": slide_height, "slides": slides}


def main() -> int:
    if len(sys.argv) != 2:
        raise SystemExit("usage: parse_document.py <controlled-file-path>")
    path = Path(sys.argv[1]).resolve(strict=True)
    suffix = path.suffix.lower()
    if suffix in {".docx", ".pptx"}:
        validate_archive(path)
    if suffix == ".md":
        lines, headings, warnings, details = parse_markdown(path)
    elif suffix == ".docx":
        lines, headings, warnings, details = parse_docx(path)
    elif suffix == ".pdf":
        lines, headings, warnings, details = parse_pdf(path)
    elif suffix == ".pptx":
        lines, headings, warnings, details = parse_pptx(path)
    else:
        raise ValueError(f"unsupported document type: {suffix}")
    text = "\n".join(lines)
    text_limit = env_limit("MAX_DOCUMENT_TEXT_CHARS", 2_000_000)
    if len(text) > text_limit:
        raise ValueError(
            f"parsed document contains {len(text)} characters; limit is {text_limit}"
        )
    result = {
        "text": text,
        "structure": {"headings": headings[:200], "lineCount": len(lines), "characterCount": len(text), **details},
        "facts": facts_from_lines(lines),
        "warnings": warnings[:100],
    }
    sys.stdout.write(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
